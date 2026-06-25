#!/usr/bin/env python3
"""Import the Google Docs DOCX export into Markdown, CSV, and assets."""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


ACTION_IDS = {
    "Detection": "detect",
    "Mitigation": "mitigate",
    "Prevention": "prevent",
}

ACTION_DESCRIPTIONS = {
    "detect": "Identify potential incidents.",
    "mitigate": "Contain an incident and restore secure operations.",
    "prevent": "Use DNS-specific steps to make recurrence less likely.",
}

YES_RID = "rId44"
NO_RID = "rId45"


def slugify(value: str) -> str:
    value = value.lower()
    value = value.replace("&", " and ")
    value = re.sub(r"[^a-z0-9]+", "-", value)
    return value.strip("-")


def clean_text(value: str) -> str:
    value = value.replace("\xa0", " ")
    value = value.replace("\u2028", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n\s+", "\n", value)
    return value.strip()


def clean_inline_text(value: str) -> str:
    value = value.replace("\xa0", " ")
    value = value.replace("\u2028", " ")
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def clean_note(value: str) -> str:
    value = clean_inline_text(value)
    value = value.strip("/")
    value = value.strip()
    if value.startswith("(") and value.endswith(")"):
        value = value[1:-1].strip()
    return value


def iter_blocks(doc: Document):
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def cell_capability(cell) -> str:
    embeds = []
    for blip in cell._tc.xpath('.//*[local-name()="blip"]'):
        rid = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
        if rid:
            embeds.append(rid)
    if YES_RID in embeds:
        return "yes"
    if NO_RID in embeds:
        return "no"
    if clean_text(cell.text).startswith("N/A"):
        return "no"
    raise ValueError(f"Could not find yes/no marker in cell: {cell.text!r}")


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def parse_definition(text: str) -> tuple[str, str]:
    for separator in (" – ", " - "):
        if separator in text:
            name, description = text.split(separator, 1)
            return clean_inline_text(name), clean_inline_text(description)
    return clean_inline_text(text), ""


def resolve_id(name: str, labels_by_id: dict[str, str]) -> str | None:
    direct = slugify(name)
    if direct in labels_by_id:
        return direct

    normalized_name = clean_inline_text(name).lower()
    normalized_base = re.sub(r"\s*[\(-].*$", "", normalized_name).strip()
    for item_id, label in labels_by_id.items():
        normalized_label = clean_inline_text(label).lower()
        if normalized_name.startswith(normalized_label):
            return item_id
        if normalized_label.startswith(normalized_name):
            return item_id
        if normalized_base and normalized_base == normalized_label:
            return item_id
    return None


def parse_matrix_table(table: Table) -> tuple[str, list[str], list[dict[str, str]]]:
    label = clean_inline_text(table.cell(0, 0).text)
    action = ACTION_IDS.get(label)
    if not action:
        raise ValueError(f"Unexpected matrix table label: {label!r}")

    headers = [clean_inline_text(cell.text) for cell in table.rows[0].cells]
    stakeholder_names = headers[1:]
    rows = []

    for row in table.rows[1:]:
        technique_name = clean_inline_text(row.cells[0].text)
        if not technique_name:
            continue
        technique_id = slugify(technique_name)
        for stakeholder_name, cell in zip(stakeholder_names, row.cells[1:]):
            note = clean_note(cell.text)
            rows.append(
                {
                    "action": action,
                    "technique_id": technique_id,
                    "stakeholder_id": slugify(stakeholder_name),
                    "capability": cell_capability(cell),
                    "note": note,
                }
            )

    return action, stakeholder_names, rows


def parse_version_history(table: Table) -> list[dict[str, str]]:
    rows = []
    for row in table.rows[1:]:
        cells = [clean_inline_text(cell.text) for cell in row.cells[:3]]
        if any(cells):
            rows.append({"version": cells[0], "date": cells[1], "comment": cells[2]})
    return rows


def markdown_for_paragraph(paragraph: Paragraph) -> str:
    text = clean_inline_text(paragraph.text)
    if not text:
        return ""
    if text.startswith(": The entity has ") and " : The entity lacks " in text:
        has_text, lacks_text = text.split(" : ", 1)
        text = f"Yes{has_text} No: {lacks_text}"

    style = paragraph.style.name if paragraph.style is not None else ""
    if style == "Heading 1":
        return f"# {text}"
    if style == "Heading 2":
        return f"## {text}"
    if style == "Heading 3":
        return f"### {text}"
    return text


def extract_sections(doc: Document, out_dir: Path) -> None:
    sections: dict[str, list[str]] = {}
    current_key = "00-title"
    current_lines: list[str] = []

    skip_h3 = {"Detection", "Mitigation", "Prevention"}
    skip_until_next_heading = False

    def flush() -> None:
        nonlocal current_lines
        if current_lines:
            if current_key != "00-title":
                sections[current_key] = current_lines
            current_lines = []

    for block in iter_blocks(doc):
        if isinstance(block, Table):
            skip_until_next_heading = False
            continue

        text = clean_inline_text(block.text)
        if not text:
            continue

        style = block.style.name if block.style is not None else ""
        if style == "Heading 1":
            continue
        if style == "Heading 2" and text == "FIRST DNS Abuse Special Interest Group":
            continue

        if style == "Heading 2":
            if text == "Version History":
                flush()
                current_key = "99-version-history"
                current_lines = []
                continue
            flush()
            current_key = f"{len(sections) + 1:02d}-{slugify(text)}"
            current_lines.append(f"## {text}")
            skip_until_next_heading = False
            continue

        if style == "Heading 3" and text in skip_h3:
            skip_until_next_heading = True
            continue

        if style in {"Heading 2", "Heading 3"}:
            skip_until_next_heading = False

        if skip_until_next_heading:
            continue

        line = markdown_for_paragraph(block)
        if line:
            current_lines.append(line)

    flush()

    content_dir = out_dir / "content"
    content_dir.mkdir(parents=True, exist_ok=True)
    for stale in content_dir.glob("*.md"):
        stale.unlink()

    for key, lines in sections.items():
        if key in {"00-title", "99-version-history"}:
            continue
        path = content_dir / f"{key}.md"
        path.write_text("\n\n".join(lines).strip() + "\n", encoding="utf-8")


def extract_assets(docx_path: Path, out_dir: Path) -> None:
    assets_dir = out_dir / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)
    wanted = {
        "word/media/image3.png": "header.png",
        "word/media/image4.png": "header-large.png",
        "word/media/image1.png": "capability-yes.png",
        "word/media/image2.png": "capability-no.png",
    }
    with zipfile.ZipFile(docx_path) as archive:
        for source, target in wanted.items():
            if source in archive.namelist():
                with archive.open(source) as src, (assets_dir / target).open("wb") as dst:
                    shutil.copyfileobj(src, dst)


def import_docx(docx_path: Path, out_dir: Path) -> None:
    doc = Document(str(docx_path))
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "data").mkdir(parents=True, exist_ok=True)

    extract_sections(doc, out_dir)
    extract_assets(docx_path, out_dir)

    stakeholders: dict[str, str] = {}
    techniques: dict[str, str] = {}
    matrix_rows: list[dict[str, str]] = []
    version_rows: list[dict[str, str]] = []

    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            text = clean_inline_text(block.text)
            if text:
                name, description = parse_definition(text)
                if name in ACTION_DESCRIPTIONS:
                    continue
                # Technique definitions are the paragraphs between the "Techniques"
                # heading and the "Stakeholders" heading. Rather than relying on
                # style metadata alone, tie the final list to the matrix row names
                # below and fill descriptions from matching paragraphs afterward.
            continue

        first = clean_inline_text(block.cell(0, 0).text)
        if first in ACTION_IDS:
            _, stakeholder_names, rows = parse_matrix_table(block)
            for name in stakeholder_names:
                stakeholders[slugify(name)] = name
            for row in block.rows[1:]:
                name = clean_inline_text(row.cells[0].text)
                if name:
                    techniques[slugify(name)] = name
            matrix_rows.extend(rows)
        elif first.startswith("Version"):
            version_rows = parse_version_history(block)

    technique_descriptions: dict[str, str] = {}
    stakeholder_descriptions: dict[str, str] = {}
    in_techniques = False
    in_stakeholders = False
    for block in iter_blocks(doc):
        if isinstance(block, Table):
            continue
        text = clean_inline_text(block.text)
        if not text:
            continue
        style = block.style.name if block.style is not None else ""
        if style == "Heading 3" and text == "Techniques":
            in_techniques = True
            in_stakeholders = False
            continue
        if style == "Heading 3" and text == "Stakeholders":
            in_techniques = False
            in_stakeholders = True
            continue
        if style.startswith("Heading"):
            in_techniques = False
            if text != "Stakeholders":
                in_stakeholders = False
            continue

        name, description = parse_definition(text)
        if in_techniques:
            sid = resolve_id(name, techniques)
            if sid:
                technique_descriptions[sid] = description
        elif in_stakeholders:
            sid = resolve_id(name, stakeholders)
            if sid:
                stakeholder_descriptions[sid] = description

    action_rows = [
        {"action": action, "label": label, "description": ACTION_DESCRIPTIONS[action]}
        for label, action in ACTION_IDS.items()
    ]
    stakeholder_rows = [
        {
            "stakeholder_id": sid,
            "label": stakeholders[sid],
            "description": stakeholder_descriptions.get(sid, ""),
        }
        for sid in stakeholders
    ]
    technique_rows = [
        {
            "technique_id": sid,
            "label": techniques[sid],
            "description": technique_descriptions.get(sid, ""),
        }
        for sid in techniques
    ]

    write_csv(out_dir / "data/actions.csv", ["action", "label", "description"], action_rows)
    write_csv(
        out_dir / "data/stakeholders.csv",
        ["stakeholder_id", "label", "description"],
        stakeholder_rows,
    )
    write_csv(
        out_dir / "data/techniques.csv",
        ["technique_id", "label", "description"],
        technique_rows,
    )
    write_csv(
        out_dir / "data/matrix.csv",
        ["action", "technique_id", "stakeholder_id", "capability", "note"],
        matrix_rows,
    )
    write_csv(
        out_dir / "data/version_history.csv",
        ["version", "date", "comment"],
        version_rows,
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--out", type=Path, default=Path("matrix"))
    args = parser.parse_args()

    import_docx(args.docx, args.out)


if __name__ == "__main__":
    main()
